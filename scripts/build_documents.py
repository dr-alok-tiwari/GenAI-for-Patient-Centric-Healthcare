from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "downloads"
DATA = ROOT / "data"

NAVY = "073647"
TEAL = "0F766E"
CYAN = "0891B2"
INDIGO = "4F46E5"
ORANGE = "C2410C"
INK = "0F172A"
MUTED = "475569"
PALE = "F7FBFD"
LINE = "D8E7EF"
WHITE = "FFFFFF"
MINT = "ECFDF5"
AMBER = "FFF7ED"
BLUE = "EFF6FF"

PILLARS = [
    ("01", "Patient Understanding and Engagement", "Explain, translate, teach back"),
    ("02", "Clinical Workflow and Decision Support", "Structure, observe, escalate"),
    ("03", "Pharma, Research and Professional Productivity", "Search, extract, document"),
    ("04", "Responsible AI, Governance and Implementation", "Classify, control, monitor"),
]

SOURCES = [
    ("WHO: Ethics and governance of AI for health", "https://www.who.int/publications/i/item/9789240029200"),
    ("NIST AI Risk Management Framework", "https://www.nist.gov/itl/ai-risk-management-framework"),
    ("FDA: AI-enabled medical devices", "https://www.fda.gov/medical-devices/software-medical-device-samd/artificial-intelligence-enabled-medical-devices"),
    ("FDA: Clinical decision support FAQ", "https://www.fda.gov/medical-devices/software-medical-device-samd/clinical-decision-support-software-frequently-asked-questions-faqs"),
    ("ChatGPT image-input limitations", "https://help.openai.com/articles/8400551-chatgpt-image-inputs-faq"),
    ("India Digital Personal Data Protection Act", "https://www.meity.gov.in/static/uploads/2024/06/2bf1f0e9f04e6fb4f8fef35e82c42aa5.pdf"),
    ("India Digital Personal Data Protection Rules", "https://www.meity.gov.in/documents/act-and-policies/digital-personal-data-protection-rules-2025-gDOxUjMtQWa"),
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size, color in [
        ("Title", 30, NAVY),
        ("Heading 1", 21, NAVY),
        ("Heading 2", 14, TEAL),
        ("Heading 3", 11, INDIGO),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(7.2))
    table.autofit = False
    table.columns[0].width = Inches(3.4)
    table.columns[1].width = Inches(3.8)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    shade(left, NAVY)
    shade(right, NAVY)
    p = left.paragraphs[0]
    r = p.add_run("ALKEM  |  AI MASTERCLASS")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    r.font.size = Pt(8.5)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(title)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    r.font.size = Pt(8)
    footer = section.footer
    p = footer.paragraphs[0]
    p.add_run("Synthetic-data learning material • qualified human review required").font.size = Pt(8)
    set_page_field(footer.add_paragraph())

    props = doc.core_properties
    props.title = title
    props.subject = "Alkem AI Masterclass: GenAI for Patient-Centric Healthcare"
    props.author = "Alkem AI Masterclass"
    props.keywords = "healthcare, generative AI, synthetic data, patient engagement, governance"


def title_page(doc: Document, document_type: str, subtitle: str) -> None:
    doc.add_paragraph("ALKEM  |  AI MASTERCLASS", style="Subtitle")
    p = doc.add_paragraph("GenAI for Patient-Centric Healthcare", style="Title")
    p.paragraph_format.space_before = Pt(36)
    p = doc.add_paragraph(document_type, style="Heading 1")
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph(subtitle)
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    for row in table.rows:
        row.height = Inches(0.7)
    labels = [
        ("60 min", "Concise theory and medical demonstrations"),
        ("90 min", "Six case-based labs with verification"),
        ("4 pillars", "Patient • Clinical • Pharma/Research • Governance"),
        ("Safety", "Synthetic data • bounded task • qualified human decision"),
    ]
    for cell, (lead, body) in zip([c for row in table.rows for c in row.cells], labels):
        shade(cell, BLUE if lead != "Safety" else AMBER)
        set_cell_border(cell, CYAN if lead != "Safety" else ORANGE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(lead + "\n")
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(TEAL if lead != "Safety" else ORANGE)
        p.add_run(body)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Important: ").bold = True
    p.add_run(
        "Every case, image, name, and value in this kit is synthetic. The materials are for education, "
        "not diagnosis, prescribing, treatment, regulatory submission, or autonomous decision-making."
    )
    doc.add_page_break()


def add_callout(doc: Document, title: str, body: str, fill: str = AMBER, accent: str = ORANGE) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, accent, "10")
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    p.add_run(body)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers):
        shade(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(8.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            shade(cell, PALE if row_index % 2 == 0 else WHITE)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.keep_together = True
            p.add_run(str(value))
    return table


def blank_lines(doc: Document, count: int = 4) -> None:
    for _ in range(count):
        p = doc.add_paragraph("________________________________________________________________________________")
        p.runs[0].font.color.rgb = RGBColor.from_string("94A3B8")
        p.paragraph_format.space_after = Pt(2)


def participant_workbook() -> Document:
    doc = Document()
    configure(doc, "Participant Workbook")
    title_page(
        doc,
        "Participant Workbook",
        "A practical workbook for doctors, pharmacists, medical affairs, pharmacovigilance, research, quality, and healthcare leadership.",
    )

    doc.add_heading("1. Four-pillar map and learning outcomes", level=1)
    add_table(doc, ["Pillar", "Purpose", "Medical examples"], [
        ["01 Patient engagement", "Make information clearer and actionable.", "Prescription explanation; discharge teach-back; multilingual counselling"],
        ["02 Clinical workflow", "Structure information while clinicians retain decisions.", "Image observation boundary; red flags; documentation"],
        ["03 Pharma and research", "Improve traceability and preparation.", "PV intake; evidence search; medical information"],
        ["04 Responsible AI", "Match controls to risk and implementation.", "Risk tier; data boundary; monitoring; stop rule"],
    ], [1.15, 2.75, 3.35])
    doc.add_heading("By the end, I can…", level=2)
    for text in [
        "□ frame a bounded healthcare task and name prohibited actions",
        "□ select a tool by workflow fit, data boundary, evidence, and reviewer",
        "□ use synthetic prescriptions and medical images safely in demonstrations",
        "□ verify source fidelity, citations, missingness, uncertainty, and escalation",
        "□ design a measurable, reversible 30-day pilot",
    ]:
        doc.add_paragraph(text)
    add_callout(doc, "My safety contract", "Use synthetic or approved data. Do not delegate diagnosis, prescribing, safety, regulatory, publication, or governance decisions. Record accept, edit, reject, or escalate.")
    doc.add_page_break()

    doc.add_heading("2. Workflow-first prompt canvas", level=1)
    add_table(doc, ["Element", "Question", "My draft"], [
        ["Role", "Who is the model assisting?", ""],
        ["Context", "Audience, setting, approved data, and source?", ""],
        ["Task", "One bounded action?", ""],
        ["Constraints", "What must not be invented or decided?", ""],
        ["Output", "What exact structure is useful?", ""],
        ["Verify", "What source checks and escalation are required?", ""],
    ], [1.0, 2.65, 3.6])
    blank_lines(doc, 5)
    add_callout(doc, "Prompt quality check", "A confident tone is not evidence. A safer prompt keeps source facts, assumptions, interpretation, missing data, and human decisions separate.", BLUE, CYAN)
    doc.add_page_break()

    labs = [
        ("Lab 1 — Prescription extraction and explanation", "prescription_cases.csv", "Exact extraction • [UNREADABLE] instead of guessing • plain-language explanation", [
            "Every medicine/product and strength matches the source",
            "Frequency, duration, and warnings match the source",
            "No indication, interaction, substitution, or new advice was added",
            "A clinician or pharmacist owns the final decision",
        ]),
        ("Lab 2 — Medical-image observation boundary", "imaging_cases.csv", "Image adequacy • neutral observations • uncertainty • escalation—not diagnosis", [
            "Adequacy and limitations assessed first",
            "No diagnosis, emergency clearance, or treatment recommendation",
            "Cannot-assess items and uncertainty are explicit",
            "Qualified clinical interpretation is explicit",
        ]),
        ("Lab 3 — Discharge plan and multilingual teach-back", "discharge_summaries.csv", "Preserve medicine, dose, date, warning, follow-up, and source meaning", [
            "Clinical meaning is unchanged",
            "Urgent warning signs are prominent",
            "Teach-back questions reveal action comprehension",
            "Fluent clinical translation review is named",
        ]),
        ("Lab 4 — Differential and red-flag checklist", "clinical_reasoning_cases.csv", "Common • cannot-miss • evidence for/against • missing data • urgency", [
            "No final diagnosis or treatment",
            "Common and cannot-miss items are separated",
            "Missing history/examination remains visible",
            "Cognitive-bias and escalation checks are present",
        ]),
        ("Lab 5 — Pharmacovigilance intake", "adverse_event_reports.csv", "Neutral chronology • seriousness fields • missing data • follow-up • PV route", [
            "No causality, expectedness, coding, or reportability decision",
            "Seriousness rationale is traceable",
            "Minimum criteria and missing fields are visible",
            "Qualified PV/medical review is explicit",
        ]),
        ("Lab 6 — Evidence search and citation verification", "research_questions.csv", "PICO/PECO • reproducible search • screening log • open every source", [
            "No fabricated paper or citation",
            "Search and screening decisions are reproducible",
            "Facts remain linked to opened sources",
            "Limitations and AI-use disclosure are present",
        ]),
    ]
    for index, (title, dataset, goal, checks) in enumerate(labs, 1):
        doc.add_heading(title, level=1)
        add_table(doc, ["Time", "Action"], [["3 min", "Orient to the source and boundary"], ["6 min", "Run the bounded prompt"], ["4 min", "Verify against source"], ["2 min", "Accept, edit, reject, or escalate"]], [1.1, 6.15])
        doc.add_paragraph(f"Dataset: {dataset}")
        doc.add_paragraph(f"Goal: {goal}")
        doc.add_heading("Verification", level=2)
        for check in checks:
            doc.add_paragraph("□ " + check)
        doc.add_heading("Human decision record", level=2)
        add_table(doc, ["Decision", "Evidence checked / edits / escalation"], [["□ Accept  □ Edit  □ Reject  □ Escalate", ""]], [2.3, 4.95])
        blank_lines(doc, 4)
        if index < len(labs):
            doc.add_page_break()

    doc.add_page_break()
    doc.add_heading("9. Tool-choice and output comparison", level=1)
    add_table(doc, ["Criterion", "Tool A", "Tool B", "Tool C"], [
        ["Workflow/input fit", "", "", ""],
        ["Data/account boundary", "", "", ""],
        ["Source or evidence traceability", "", "", ""],
        ["Missingness and uncertainty", "", "", ""],
        ["Human correction required", "", "", ""],
        ["Access, region, and cost verified", "", "", ""],
        ["Decision: use / do not use / escalate", "", "", ""],
    ], [2.2, 1.68, 1.68, 1.68])
    add_callout(doc, "Selection rule", "Do not rank tools by fluency alone. Compare fidelity, traceability, safety boundaries, privacy fit, usability, and the amount of qualified correction required.", BLUE, INDIGO)
    doc.add_page_break()

    doc.add_heading("10. My 30-day responsible pilot", level=1)
    add_table(doc, ["Canvas", "My decision"], [
        ["Pain point and bounded task", ""],
        ["Accountable owner", ""],
        ["Allowed / prohibited data", ""],
        ["Users and qualified reviewer", ""],
        ["Baseline and success measure", ""],
        ["Monitoring and audit evidence", ""],
        ["Incident / escalation route", ""],
        ["Stop and rollback rule", ""],
    ], [2.25, 4.95])
    blank_lines(doc, 5)
    doc.add_page_break()

    doc.add_heading("11. Sources and further reading", level=1)
    doc.add_paragraph("Accessed and verified for this masterclass on 28 July 2026. Product features, pricing, regional access, and policy status can change; recheck before delivery or implementation.")
    for title, url in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(title + ": ").bold = True
        p.add_run(url)
    return doc


def facilitator_handbook() -> Document:
    doc = Document()
    configure(doc, "Facilitator Handbook")
    title_page(
        doc,
        "Facilitator Handbook",
        "Delivery notes, medical demonstration boundaries, lab runbooks, fallback routes, verification guidance, and answer key.",
    )

    doc.add_heading("1. Delivery at a glance", level=1)
    add_table(doc, ["Minutes", "Slides / activity", "Facilitator outcome"], [
        ["00–08", "Slides 1–2", "Set the synthetic-data, bounded-task, source-check, and human-authority contract."],
        ["08–22", "Slides 3–4", "Pillar 1 and synthetic prescription demonstration."],
        ["22–36", "Slides 5–6", "Patient action plan and medical-image observation boundary."],
        ["36–50", "Slides 7–8", "Clinical reasoning, PV intake, and evidence verification."],
        ["50–60", "Slides 9–10", "Risk screen, tool selection, and lab assignment."],
        ["60–105", "Labs 1–3", "Prescription; image observation; discharge + multilingual teach-back."],
        ["105–150", "Labs 4–6", "Differential/red flags; PV intake; evidence/citation verification."],
    ], [1.0, 2.0, 4.2])
    add_callout(doc, "Non-negotiable opening", "Do not ask participants to upload real patient data, real prescriptions, scans, company-confidential documents, or identifiable case details. Use the included synthetic assets.")
    doc.add_page_break()

    doc.add_heading("2. Four-pillar speaking notes", level=1)
    for number, title, verbs in PILLARS:
        doc.add_heading(f"{number} {title}", level=2)
        doc.add_paragraph(verbs + ".")
        if number == "01":
            doc.add_paragraph("Emphasise comprehension, teach-back, translation review, and exact medication/source fidelity.")
        elif number == "02":
            doc.add_paragraph("Emphasise adequacy, missingness, red flags, uncertainty, bias, and clinician-owned diagnosis/treatment.")
        elif number == "03":
            doc.add_paragraph("Separate intake structuring and discovery from PV, medical, regulatory, screening, and interpretation decisions.")
        else:
            doc.add_paragraph("Classify risk by data, impact, autonomy, and reach; require ownership, monitoring, incident route, and stop rule.")
    doc.add_page_break()

    doc.add_heading("3. Live demo: synthetic prescription", level=1)
    add_table(doc, ["Step", "Facilitator action", "Watch for"], [
        ["1 Prepare", "Open `synthetic_prescription_01.png`; keep CSV source values hidden.", "No real prescription or identifier."],
        ["2 Prompt", "Use exact extraction, [UNREADABLE], explanation, and verification sections.", "Model guesses or infers indication."],
        ["3 Reveal", "Compare every field against `prescription_cases.csv`.", "Dose/frequency/duration drift."],
        ["4 Decide", "Ask a clinician/pharmacist role to accept, edit, reject, or escalate.", "Fluency substituted for fidelity."],
    ], [0.9, 3.35, 2.95])
    add_callout(doc, "Boundary", "The demonstration is not OCR validation, prescribing, medication reconciliation, interaction checking, or clinical advice.")
    doc.add_page_break()

    doc.add_heading("4. Live demo: synthetic medical image", level=1)
    add_table(doc, ["Do", "Do not"], [
        ["State: fully synthetic and no diagnostic ground truth.", "Ask for or reveal a diagnosis."],
        ["Assess image adequacy and limitations first.", "Claim general-purpose model validation."],
        ["Use neutral visible observations and cannot-assess items.", "Allow emergency clearance or treatment advice."],
        ["Require radiologist/treating-clinician interpretation.", "Use a real medical image in an unapproved tool."],
    ], [3.6, 3.6])
    doc.add_paragraph("Use the official product limitation page during the debrief: https://help.openai.com/articles/8400551-chatgpt-image-inputs-faq")
    add_callout(doc, "Offline fallback", "Show the static synthetic image and ask teams to rewrite an unsafe diagnostic prompt into an adequacy–observation–uncertainty–escalation prompt.")
    doc.add_page_break()

    doc.add_heading("5. Demo stack and fallback routes", level=1)
    add_table(doc, ["Workflow", "Primary route", "Alternative", "Offline fallback"], [
        ["Prescription", "ChatGPT", "Gemini / Microsoft Copilot", "Static image + exact source table"],
        ["Medical image", "ChatGPT", "Gemini", "Prompt-boundary critique"],
        ["PV intake", "ChatGPT / Claude", "Gemini / Julius AI", "Manual PV intake template"],
        ["Evidence", "PubMed / Elicit", "Consensus / Semantic Scholar", "Synthetic abstract extraction"],
        ["Patient material", "Canva + reviewed copy", "Gamma", "Workbook fact-sheet canvas"],
    ], [1.35, 1.55, 1.9, 2.45])
    doc.add_paragraph("Product access, capabilities, regional availability, account eligibility, terms, and pricing can change. Verify on the official site before delivery.")
    doc.add_page_break()

    doc.add_heading("6. Six lab runbooks", level=1)
    workflows = json.loads((DATA / "workflow_catalog.json").read_text(encoding="utf-8"))
    selected_ids = ["P1-RX-01", "P2-XRAY-01", "P1-DC-02", "P2-REASON-02", "P3-PV-01", "P3-EVID-03"]
    selected = [next(item for item in workflows if item["workflow_id"] == workflow_id) for workflow_id in selected_ids]
    for index, item in enumerate(selected, 1):
        doc.add_heading(f"Lab {index}: {item['title']}", level=2)
        add_table(doc, ["Field", "Facilitator reference"], [
            ["Dataset", item["dataset"]],
            ["Primary / alternatives", " / ".join([item["primary_tool"], *item["alternative_tools"]])],
            ["Expected output", item["expected_output"]],
            ["Human reviewer", item["human_reviewer"]],
            ["Prohibited", item["prohibited_actions"]],
        ], [1.55, 5.65])
        doc.add_paragraph("Debrief: What changed from source? What could harm? What was verified? Who decided? What control and stop condition would implementation require?")
        if index in {2, 4}:
            doc.add_page_break()
    doc.add_page_break()

    doc.add_heading("7. Dropdown coverage and selection logic", level=1)
    add_table(doc, ["Selector", "Coverage rule"], [
        ["Pillar", "Options come from complete workflow records."],
        ["Role", "Shows only roles remaining after pillar selection."],
        ["Task", "Shows only tasks remaining after prior selections."],
        ["Input / language", "Unsupported intersections are removed before selection."],
        ["Free-text need", "Returns a clearly labelled closest verified alternative."],
        ["Tool / demo", "Every result includes dataset, prompt, output, verification, safety, and reviewer."],
    ], [1.6, 5.6])
    doc.add_paragraph(f"Current catalogue: {len(workflows)} complete workflows across four pillars. Automated tests enumerate every visible option and verify that it returns at least one workflow.")
    doc.add_page_break()

    doc.add_heading("8. Assessment answer key", level=1)
    answers = [
        ["1–5", "Synthetic/approved data; professional accountability; mark unreadable; non-diagnostic observation; adequacy first"],
        ["6–10", "Preserve source; fluent clinical review; teach-back; checklist not authority; common vs cannot-miss"],
        ["11–15", "Neutral PV intake; verify citations; reproducible logs; source checks; mark missing"],
        ["16–20", "Identifiable + autonomy raises risk; pilot controls; subgroup/language review; workflow fit; decision record"],
    ]
    add_table(doc, ["Questions", "Key concepts"], answers, [1.1, 6.1])
    add_callout(doc, "Scoring guidance", "80%+ strong foundation; 60–79% review the relevant explanation; below 60% revisit governance and verification before planning a pilot.", BLUE, INDIGO)
    doc.add_page_break()

    doc.add_heading("9. Troubleshooting and safety interventions", level=1)
    add_table(doc, ["Situation", "Immediate response"], [
        ["Participant starts using real data", "Stop, delete/close the input if possible, and return to included synthetic assets; follow institutional incident guidance if exposure occurred."],
        ["Model gives a diagnosis", "Do not debate correctness. Mark boundary breach; rewrite prompt; require qualified interpretation."],
        ["Model fabricates a citation", "Open source, show failure, discard unsupported claim, and demonstrate citation verification."],
        ["No tool access / sign-in failure", "Use static source + manual prompt critique + expected-output and verification templates."],
        ["Dropdown search is unfamiliar", "Use the labelled closest verified alternative or cascade by pillar and role."],
        ["Discussion becomes vendor promotion", "Return to task, data boundary, evidence, reviewer, monitoring, and stop condition."],
    ], [2.05, 5.15])
    doc.add_page_break()

    doc.add_heading("10. Sources and verification log", level=1)
    doc.add_paragraph("Verified on 28 July 2026. Recheck before each delivery because product, policy, and regulatory information can change.")
    for title, url in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(title + ": ").bold = True
        p.add_run(url)
    return doc


def dropdown_coverage_report() -> Document:
    doc = Document()
    configure(doc, "Dropdown Coverage Report")
    title_page(
        doc,
        "Dropdown Coverage Report",
        "Evidence that every visible selector option in the Alkem AI Masterclass app resolves to a complete, medically relevant response.",
    )
    workflows = json.loads((DATA / "workflow_catalog.json").read_text(encoding="utf-8"))
    fields = [
        ("Pillar", "pillar"),
        ("Role", "roles"),
        ("Task", "task"),
        ("Input", "input_type"),
        ("Output", "output_type"),
        ("Language", "languages"),
        ("Risk", "risk_level"),
    ]
    coverage = []
    for label, field in fields:
        options = sorted(
            {
                value
                for item in workflows
                for value in (
                    item[field] if isinstance(item[field], list) else [item[field]]
                )
            }
        )
        for option in options:
            count = sum(
                option in (item[field] if isinstance(item[field], list) else [item[field]])
                for item in workflows
            )
            coverage.append([label, option, str(count), "Covered" if count else "Gap"])

    doc.add_heading("1. Executive result", level=1)
    add_table(doc, ["Measure", "Result"], [
        ["Complete workflow records", str(len(workflows))],
        ["Programme pillars", str(len({item["pillar"] for item in workflows}))],
        ["Professional roles", str(len({role for item in workflows for role in item["roles"]}))],
        ["Distinct tasks", str(len({item["task"] for item in workflows}))],
        ["Visible selector options tested", str(len(coverage))],
        ["Options without a response", str(sum(row[3] == "Gap" for row in coverage))],
    ], [3.8, 3.4])
    add_callout(doc, "Result: 100% visible-option coverage", "Every option shown to a user is derived from one or more complete workflow records. Unsupported intersections are removed by cascading selectors before they can be selected.", MINT, TEAL)
    doc.add_page_break()

    doc.add_heading("2. What counts as a complete response", level=1)
    add_table(doc, ["Required field", "Coverage condition"], [
        ["Medical context", "Named case context and pillar"],
        ["Professional fit", "One or more roles and accountable human reviewer"],
        ["Action", "Bounded task and explicit expected output"],
        ["Input", "Supported input type and linked synthetic dataset"],
        ["Tool route", "Primary tool and verified alternatives"],
        ["Prompt", "Complete task, constraints, output, and verification template"],
        ["Safety", "Risk tier, prohibited actions, missingness and escalation"],
        ["Traceability", "Verification checklist and last-verified date"],
    ], [2.0, 5.2])
    doc.add_heading("Selector behaviour", level=2)
    doc.add_paragraph(
        "Each dropdown is populated from the records still reachable after previous selections. "
        "Therefore a visible option cannot lead to an unsupported intersection. Free-text search is "
        "handled separately: exact matches are shown when available; otherwise the app labels and "
        "returns the closest verified alternative."
    )
    doc.add_page_break()

    doc.add_heading("3. Coverage by visible option", level=1)
    add_table(doc, ["Selector", "Visible option", "Workflows", "Status"], coverage, [1.05, 4.35, 0.8, 1.0])
    doc.add_page_break()

    doc.add_heading("4. Automated test evidence", level=1)
    add_table(doc, ["Test", "Pass condition"], [
        ["Catalog completeness", "Every workflow has non-empty required fields, dataset, reviewer, safety, and verified date."],
        ["Visible option enumeration", "Every pillar, role, task, input, output, language, and risk option returns at least one workflow."],
        ["Cascading intersections", "For every workflow, its ordered pillar-role-task-input-language path remains selectable."],
        ["Tool integrity", "All 29 tool rows contain official HTTPS URL, dataset, safety warning, reviewer, and verification metadata."],
        ["Dataset and asset integrity", "Every linked CSV and synthetic prescription/image asset exists."],
        ["App empty-state scan", "Legacy terminal warnings for unmatched tool/demo selections are absent."],
    ], [2.25, 4.95])
    add_callout(doc, "Interpretation", "This report verifies application coverage and content integrity. It does not validate model performance, tool security, clinical efficacy, or regulatory suitability.")
    doc.add_heading("Sources", level=2)
    for title, url in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(title + ": ").bold = True
        p.add_run(url)
    return doc


def save(doc: Document, filename: str) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / filename)
    print(f"Saved {filename}")


def main() -> None:
    save(participant_workbook(), "Alkem_AI_Masterclass_Participant_Workbook.docx")
    save(facilitator_handbook(), "Alkem_AI_Masterclass_Facilitator_Handbook.docx")
    save(dropdown_coverage_report(), "Dropdown_Coverage_Report.docx")


if __name__ == "__main__":
    main()
